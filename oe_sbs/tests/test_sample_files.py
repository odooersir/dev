#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Regression harness for the EXTRACTION layer, run against real supplier files.

WHY A SECOND HARNESS
--------------------
test_parsing_cases.py feeds strings straight to the parsers. That catches
parser bugs, but it cannot catch the ones that happen BEFORE the parser is
reached - and several of the worst bugs were exactly that:

  * a conditions sidebar counted as a table column, so the whole block was
    never scanned (Rhode, MinMaxDeals)
  * a footer written in the EAN column pushed the end of the table past it, so
    the footer never qualified as a footer (Beauty Care Global)
  * a multi-line cell split into fragments that all shared one (row, col), so
    the label/value lookup picked the LAST line (Biologique Recherche)
  * a label in column B whose value sat in column C, where only the label's
    address was recorded (Supernova)

In each case the string the parser eventually saw was the wrong string, and a
string-level test happily passed. This harness starts from the .xlsx.

    python3 oe_sbs/tests/test_sample_files.py            # all files
    python3 oe_sbs/tests/test_sample_files.py Rhode      # just matching ones

No Odoo, no database: the pure functions are loaded out of import_wizard.py by
source. Anything that needs res.country / res.currency (incoterm place
resolution, currency conversion, margins) is deliberately NOT asserted here -
those belong in an Odoo test with a real database.
"""
import json
import logging
import os
import re
import sys
import calendar
import math
import textwrap

logging.getLogger('sbs.harness').addHandler(logging.NullHandler())
logging.getLogger('sbs.harness').propagate = False

HERE = os.path.dirname(os.path.abspath(__file__))
SAMPLES = os.path.join(HERE, 'samples')
WIZARD = os.path.join(HERE, '..', 'wizard', 'import_wizard.py')
EXPECTED = os.path.join(SAMPLES, 'expected.json')

try:
    from openpyxl import load_workbook
except ImportError:
    raise SystemExit('openpyxl is required: pip install openpyxl')


def load_namespace():
    src = open(WIZARD, encoding='utf-8').read()
    # the wizard logs freely; give it a sink instead of stubbing every call
    ns = {'re': re, 'math': math, 'calendar': calendar, 'date': __import__('datetime').date,
          '_logger': logging.getLogger('sbs.harness'),
          '_': lambda text, *a: (text % a) if a else text}

    def grab(pattern, wrap_class=None):
        m = re.search(pattern, src, re.S | re.M)
        if not m:
            raise SystemExit('harness out of date: %r not found' % pattern)
        code = m.group(0)
        if wrap_class:
            code = 'class %s:\n' % wrap_class + code
            exec(compile(code, '<wizard>', 'exec'), ns)
        else:
            exec(compile(textwrap.dedent(code), '<wizard>', 'exec'), ns)

    grab(r'_MONTH_NAMES = \{.*?(?=\ndef standardize_lead_time)')
    for fn in ('col_letter_to_index', 'index_to_col_letter', '_normalize',
               '_looks_like_label', '_inline_value', '_validate',
               'standardize_lead_time', 'standardize_mov', '_is_take_all',
               '_is_monetary', '_norm_country'):
        grab(r'def %s\b.*?(?=\ndef |\n@|\nclass |\n_[A-Z])' % fn)
    grab(r'_TAKE_ALL_RE = re\.compile\(.*?(?=\ndef _is_take_all)')
    grab(r'_SHELF_LIFE_HINTS = re\.compile\(.*?(?=\ndef _incoterm_context_ok)')
    grab(r'def _incoterm_context_ok.*?(?=\ndef |\n@)')
    grab(r'_META_KEYWORDS = .*?(?=\ndef _looks_like_metadata_statement)')
    grab(r'def _looks_like_metadata_statement.*?(?=\ndef _validate)')
    grab(r'^FB_NONE\s+=.*?^SCOPE_BOTH.*?$')
    grab(r'DYNAMIC_FIELDS = \{.*?\n\}\n')
    grab(r'def extract_metadata_cells.*?(?=\ndef pre_ai_extract)')
    grab(r'def pre_ai_extract.*?(?=\ndef |\nclass |\n@)')
    grab(r'    # Wording that means .*?(?=\n    def _take_all_mov)', wrap_class='_W')
    grab(r'    def _mov_to_number\(self.*?(?=\n    @staticmethod\n    def _round_up_to)',
         wrap_class='_MovHolder')
    return ns


NS = load_namespace()


class _Helpers(NS['_MovHolder'], NS['_W']):
    @staticmethod
    def _round_up_to(value, step):
        return int(math.ceil(float(value) / step) * step)

    def _get_default_profit_margin(self):
        return 0.0        # the harness asserts raw amounts, never grossed up


H = _Helpers()
# Read straight from DYNAMIC_FIELDS instead of restating the patterns here:
# a hardcoded copy drifts, and a harness that tests yesterday's regex is
# worse than no harness ('EU + CLEAN' passed in the wizard and failed here).
T1_RE = re.compile(NS['DYNAMIC_FIELDS']['t1']['anchor_pattern'], re.IGNORECASE)
T2_RE = re.compile(NS['DYNAMIC_FIELDS']['t2']['anchor_pattern'], re.IGNORECASE)
E1_RE = re.compile(NS['DYNAMIC_FIELDS']['euro1']['anchor_pattern'], re.IGNORECASE)


class FakeTemplate(object):
    """Stands in for sbs.import.template.

    The mappings come from expected.json, which is meant to be filled from the
    real thing: Import Templates -> Export Mapping (JSON). Guessed mappings make
    the test measure a file that nobody actually imports that way.
    """

    def __init__(self, mapping):
        self.__dict__.update(mapping)
        self._fields = list(mapping)

    def __getattr__(self, name):
        return False


def offer_fields(path, mapping):
    """Run the extraction chain and return the offer-level values it found."""
    worksheet = load_workbook(path, data_only=True).active
    sheet = [[c.value for c in row] for row in worksheet.iter_rows()]

    cells = NS['extract_metadata_cells'](sheet, FakeTemplate(mapping))
    # 'both' counts too - lead time is scoped BOTH (a column when the file has
    # one, an offer-level statement otherwise) and leaving it out made every
    # lead time in this harness read as None.
    offer_cfg = {k: v for k, v in NS['DYNAMIC_FIELDS'].items()
                 if v.get('scope') in ('offer', 'both')}
    found = NS['pre_ai_extract'](cells, offer_cfg)

    def raw(key):
        return (found.get(key) or {}).get('value')

    # Mirror the ONE resolution rule the wizard applies between extraction and
    # storage: an MOQ expressed in money is really the minimum order VALUE, so
    # it moves to MOV. Without this the harness reported a monetary MOQ as an
    # MOQ and an empty MOV, which is not what gets written to sbs.data.
    moq_v = raw('moq')
    mov_v = raw('mov')
    if moq_v and NS['_is_monetary'](moq_v):
        mov_v, moq_v = moq_v, None

    return {
        'metadata_cells': ['%s|%s' % (c['addr'], c['value']) for c in cells],
        'lead_time': NS['standardize_lead_time'](raw('lead_time') or ''),
        'mov_amount': H._mov_to_number(mov_v) if mov_v else None,
        'moq_raw': moq_v,
        'incoterm_raw': raw('incoterms'),
        'payment_term': raw('payment_term'),
        't1': H._resolve_yes_no(raw('t1'), T1_RE),
        't2': H._resolve_yes_no(raw('t2'), T2_RE, allow_available=True),
        'euro1': H._resolve_yes_no(raw('euro1'), E1_RE),
    }


def run(filter_text=None):
    if not os.path.exists(EXPECTED):
        raise SystemExit('missing %s' % EXPECTED)
    expected = json.load(open(EXPECTED, encoding='utf-8'))

    checked = failed = skipped = 0
    for filename in sorted(expected):
        if filter_text and filter_text.lower() not in filename.lower():
            continue
        spec = expected[filename]
        path = os.path.join(SAMPLES, filename)
        if not os.path.exists(path):
            print('SKIP  %s (file not in samples/)' % filename)
            skipped += 1
            continue
        if not spec.get('mapping'):
            print('SKIP  %s (no mapping yet - export it from the template)'
                  % filename)
            skipped += 1
            continue

        try:
            got = offer_fields(path, spec['mapping'])
        except Exception as exc:            # a crash is a failure, not a skip
            print('FAIL  %s raised %s: %s' % (filename, type(exc).__name__, exc))
            failed += 1
            continue

        print('\n-- %s' % filename)
        for key, want in sorted(spec.get('expect', {}).items()):
            checked += 1
            actual = got.get(key)
            if key == 'metadata_contains':
                ok = all(any(frag in cell for cell in got['metadata_cells'])
                         for frag in want)
                actual = want if ok else got['metadata_cells'][:8]
            else:
                ok = actual == want
            if ok:
                print('   ok   %-14s %r' % (key, actual))
            else:
                failed += 1
                print('   FAIL %-14s got %r  want %r' % (key, actual, want))

    print('\n%s checks, %s failed, %s file(s) skipped'
          % (checked, failed, skipped))
    return 1 if failed else 0


if __name__ == '__main__':
    sys.exit(run(sys.argv[1] if len(sys.argv) > 1 else None))
